#!/usr/bin/env python3
"""
CONFIDENCIAL: no se publica.

Arma las planillas de revisión en Excel (/data/revision/*.xlsx) a partir de data/revision/planillas.json,
que genera scripts/generar_revision.js. Cada bloque (Revisión completa, Muestreo, Resto) va en una hoja.

Solo se pueden editar "Aprobado" (desplegable sí/no), "Revisó" (desplegable G/A) y "Comentario"; el resto de la hoja
está protegido (sin contraseña) y en gris. La columna "Huella" queda oculta: scripts/leer_completadas.py la usa para
comprobar que la respuesta no cambió desde que se generó la planilla.

Uso: python3 scripts/planillas_xlsx.py   (lo ejecuta generar_revision.js al final)
Requiere openpyxl.
"""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVISION = os.path.join(RAIZ, 'data', 'revision')

# (encabezado, clave en planillas.json, ancho, editable)
COLUMNAS = [
    ('País', 'nombre', 16, False), ('ISO', 'pais', 6, False), ('#', 'pregunta', 7, False), ('Pregunta', 'texto_pregunta', 40, False),
    ('Respuesta', 'respuesta', 11, False), ('Dato', 'dato', 60, False), ('Dato que se publica', 'dato_publico', 40, False),
    ('Link', 'link', 30, False), ('Fuente y fecha de consulta', 'fuente', 36, False), ('Criterio aplicado', 'criterio', 50, False),
    ('Fuentes revisadas', 'fuentes_revisadas', 60, False), ('Estado', 'estado', 18, False),
    ('Aprobado', 'aprobado', 11, True), ('Revisó', 'reviso', 9, True), ('Comentario', 'comentario', 45, True),
    ('Huella', 'huella', 14, False),
]
GRIS = PatternFill('solid', fgColor='E7E7E7')
ENCABEZADO = PatternFill('solid', fgColor='000030')
EDITABLE = PatternFill('solid', fgColor='FFFFFF')
MAX_CELDA = 32000  # Excel admite hasta 32.767 caracteres por celda


def texto(f, clave):
    v = f.get(clave, '')
    if clave == 'fuentes_revisadas':
        v = '\n'.join(f"{o['nombre']} ({o['url']}): {o.get('resultado') or ''}" for o in v)
    if clave == 'respuesta' and f.get('critica'):
        v = f'{v} ⚡'
    if clave == 'link' and f.get('otros_links'):
        v = '\n'.join([v] + [o['url'] for o in f['otros_links']])
    return str(v)[:MAX_CELDA]


def hoja(wb, titulo, filas):
    ws = wb.create_sheet(titulo)
    for i, (enc, _, ancho, editable) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=1, column=i, value=enc)
        c.font = Font(bold=True, color='FEFEFE')
        c.fill = ENCABEZADO
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = ancho
    for n, f in enumerate(filas, start=2):
        for i, (_, clave, _, editable) in enumerate(COLUMNAS, start=1):
            c = ws.cell(row=n, column=i, value=texto(f, clave) or None)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if editable:
                c.fill = EDITABLE
                c.protection = Protection(locked=False)
            else:
                c.fill = GRIS
                c.protection = Protection(locked=True)
            if clave == 'link' and f.get('link'):
                c.hyperlink = f['link']
                c.font = Font(color='0059FF', underline='single')
    ultima = max(len(filas) + 1, 2)
    col = {clave: ws.cell(row=1, column=i).column_letter for i, (_, clave, _, _) in enumerate(COLUMNAS, start=1)}
    si_no = DataValidation(type='list', formula1='"sí,no"', allow_blank=True, showErrorMessage=True,
                           errorTitle='Aprobado', error='Elegí "sí" o "no".')
    revisor = DataValidation(type='list', formula1='"G,A"', allow_blank=True, showErrorMessage=True,
                             errorTitle='Revisó', error='Elegí G (Giuliano) o A (Agustina).')
    ws.add_data_validation(si_no)
    ws.add_data_validation(revisor)
    si_no.add(f"{col['aprobado']}2:{col['aprobado']}{ultima}")
    revisor.add(f"{col['reviso']}2:{col['reviso']}{ultima}")
    ws.column_dimensions[col['huella']].hidden = True
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{col["comentario"]}{ultima}'
    # Protegida sin contraseña: evita cambios accidentales y deja filtrar y ajustar anchos.
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    return ws


def instrucciones(wb, pl, bloques):
    ws = wb.active
    ws.title = 'Instrucciones'
    lineas = [
        pl['titulo'],
        'Confidencial: no se publica.',
        f"Generada el {pl['generado']} con node scripts/generar_revision.js.",
        '',
        'Cómo revisar:',
        '1. Abrí el link de cada fila y verificá que el dato coincida con la fuente.',
        '2. Completá "Aprobado" (sí/no) y "Revisó" (G = Giuliano, A = Agustina). Si algo está mal o falta, escribilo en "Comentario".',
        '3. Solo se pueden editar esas tres columnas; el resto de la hoja está protegido.',
        '4. Guardá el archivo y subilo a data/revision/completadas/ en GitHub (Add file → Upload files).',
        '   Las filas con "sí" y sin comentario pasan a aprobado. Las que tienen "no" o un comentario se corrigen y se listan para confirmar.',
        '',
        'Hojas:',
    ] + [f"- {b['titulo']} ({len(pl['bloques'][b['id']])} filas): {b['nota']}" for b in bloques]
    if pl.get('nota'):
        lineas += ['', pl['nota']]
    if pl.get('puntos'):
        lineas += ['', 'Puntos a mirar:'] + [f'- {p}' for p in pl['puntos']]
    for i, t in enumerate(lineas, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if i == 1:
            c.font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 140
    ws.protection.sheet = True


def main():
    with open(os.path.join(REVISION, 'planillas.json'), encoding='utf-8') as fh:
        datos = json.load(fh)
    for pl in datos['planillas']:
        wb = Workbook()
        instrucciones(wb, pl, datos['bloques'])
        for b in datos['bloques']:
            hoja(wb, b['titulo'], pl['bloques'][b['id']])
        salida = os.path.join(REVISION, f"{pl['archivo']}.xlsx")
        wb.save(salida)
        print(f'Excel generado en data/revision/{pl["archivo"]}.xlsx')


if __name__ == '__main__':
    main()
