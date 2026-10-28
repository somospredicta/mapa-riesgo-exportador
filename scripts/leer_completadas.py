#!/usr/bin/env python3
"""
CONFIDENCIAL: no se publica.

Lee las planillas de revisión completadas que se suben a data/revision/completadas/ (.xlsx, o .md editado a mano)
y actualiza /data/respuestas/:
  - "Aprobado" = sí, con "Revisó" (G o A) y sin comentario → la respuesta pasa a "aprobado" y guarda revisado_por,
    revisado_fecha y revision_origen.
  - "Aprobado" = no, o cualquier fila con comentario → no se toca: va al informe "Para corregir", para corregirla
    según el comentario y confirmarla con los revisores.
  - Si la respuesta cambió desde que se generó la planilla (la huella no coincide), o falta quién revisó, tampoco se
    aprueba: va al informe.
Aprobación por muestreo: si todas las filas de la hoja "Muestreo" de una pregunta están aprobadas sin comentarios, las filas
de esa pregunta en la hoja "Resto" (sin decisión propia y sin cambios desde la planilla) pasan a "aprobada_por_muestreo", con
quién aprobó el muestreo y la fecha. Si alguna fila del muestreo tiene "no" o comentario, el resto de esa pregunta no se
aprueba. Las preguntas de EXCLUIDAS_MUESTREO no se aprueban por muestreo.
Cada planilla procesada se mueve a data/revision/completadas/procesadas/ y deja un informe_<planilla>_<fecha>.md.

Uso: python3 scripts/leer_completadas.py
     python3 scripts/leer_completadas.py --reaplicar   (vuelve a leer las planillas de procesadas/, por ejemplo para aplicar
                                                        el muestreo a planillas ya procesadas; lo ya aprobado no cambia)
Después: node scripts/calcular.js && node scripts/generar_revision.js
"""
import datetime
import glob
import hashlib
import json
import os
import re
import shutil
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RESPUESTAS = os.path.join(RAIZ, 'data', 'respuestas')
COMPLETADAS = os.path.join(RAIZ, 'data', 'revision', 'completadas')
PROCESADAS = os.path.join(COMPLETADAS, 'procesadas')
REVISORES = {'G': 'Giuliano', 'A': 'Agustina'}
APROBADOS = ('aprobado', 'aprobada_por_muestreo')
EXCLUIDAS_MUESTREO = {'A1'}  # A1 queda fuera del muestreo hasta terminar su corrección (sesión 6)
SI = {'sí', 'si', 's', 'yes', '✔', '✓', 'ok', 'x'}
NO = {'no', 'n', '✘', '✗'}


def huella(r):
    """Igual que en scripts/generar_revision.js."""
    partes = ['' if r.get(k) is None else str(r.get(k)) for k in ('respuesta', 'dato', 'dato_publico', 'anio_dato', 'fuente_url', 'consultado')]
    return hashlib.sha1('\u0001'.join(partes).encode('utf-8')).hexdigest()[:12]


def limpiar(v):
    return '' if v is None else str(v).strip()


def filas_xlsx(ruta):
    from openpyxl import load_workbook
    wb = load_workbook(ruta, data_only=True)
    for ws in wb.worksheets:
        if ws.title == 'Instrucciones':
            continue
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            continue
        enc = [limpiar(c) for c in filas[0]]
        if 'ISO' not in enc or '#' not in enc:
            continue
        for n, f in enumerate(filas[1:], start=2):
            d = dict(zip(enc, f))
            if not limpiar(d.get('ISO')) or not limpiar(d.get('#')):
                continue  # fila vacía
            yield {'hoja': ws.title, 'fila': n, 'iso': limpiar(d.get('ISO')), 'pregunta': limpiar(d.get('#')),
                   'aprobado': limpiar(d.get('Aprobado')), 'reviso': limpiar(d.get('Revisó')).upper(),
                   'comentario': limpiar(d.get('Comentario')), 'huella': limpiar(d.get('Huella'))}


def filas_md(ruta, nombres):
    """Tablas Markdown de las planillas: la primera columna es el país (o, en planillas viejas, una tabla por país)."""
    texto = open(ruta, encoding='utf-8').read()
    pais_actual = None
    enc = None
    for n, linea in enumerate(texto.splitlines(), start=1):
        m = re.match(r'^## .*\((\w{3})\)\s*$', linea)
        if m:
            pais_actual = m.group(1)
        if not linea.startswith('|'):
            enc = None if not linea.strip() else enc
            continue
        celdas = [c.strip() for c in re.split(r'(?<!\\)\|', linea.strip())[1:-1]]
        if celdas and celdas[0] in ('País', '#'):
            enc = celdas
            continue
        if not enc or set(''.join(celdas)) <= set('-: '):
            continue
        d = dict(zip(enc, celdas))
        iso = nombres.get(d.get('País', ''), pais_actual)
        aprobado = d.get('Aprobado', d.get('Aprobado (Agustina)', ''))
        reviso = d.get('Revisó', 'A' if 'Aprobado (Agustina)' in d else '')
        yield {'hoja': os.path.basename(ruta), 'fila': n, 'iso': iso or '', 'pregunta': d.get('#', ''), 'aprobado': aprobado,
               'reviso': reviso.upper(), 'comentario': d.get('Comentario', '').strip('— '), 'huella': ''}


def procesar(ruta, datos, nombres, hoy):
    base = os.path.basename(ruta)
    filas = list(filas_xlsx(ruta) if ruta.endswith('.xlsx') else filas_md(ruta, nombres))
    aprobadas, ya, corregir, revisar = [], [], [], []
    for f in filas:
        ap = f['aprobado'].lower()
        if not ap and not f['comentario']:
            continue
        clave = f"{f['iso']} {f['pregunta']}"
        d = datos.get(f['iso'])
        r = d['respuestas'].get(f['pregunta']) if d else None
        if r is None:
            revisar.append((f, 'No se encontró esa respuesta en /data/respuestas/.'))
            continue
        if ap in NO or f['comentario']:
            corregir.append((f, r))
            continue
        if ap not in SI:
            revisar.append((f, f'Valor de "Aprobado" no reconocido: "{f["aprobado"]}".'))
            continue
        if f['reviso'] not in REVISORES:
            revisar.append((f, 'Falta "Revisó" (G o A).'))
            continue
        if f['huella'] and f['huella'] != huella(r):
            revisar.append((f, 'La respuesta cambió desde que se generó la planilla: hay que revisarla de nuevo.'))
            continue
        if r.get('estado') in APROBADOS:
            ya.append(clave)
            continue
        r['estado'] = 'aprobado'
        r['revisado_por'] = f['reviso']
        r['revisado_fecha'] = hoy
        r['revision_origen'] = f'{base} ({f["hoja"]}, fila {f["fila"]})'
        d['_cambio'] = True
        aprobadas.append(f"{clave} ({f['reviso']})")
    muestreo, bloqueadas = aplicar_muestreo(filas, datos, base, hoy, revisar)
    return aprobadas, ya, corregir, revisar, muestreo, bloqueadas


def aplicar_muestreo(filas, datos, base, hoy, revisar):
    """Aprueba por muestreo las filas de "Resto" de cada pregunta cuyo muestreo completo quedó aprobado sin comentarios."""
    por_pregunta = {}
    for f in filas:
        if f['hoja'] == 'Muestreo':
            por_pregunta.setdefault(f['pregunta'], []).append(f)
    aprobadas, bloqueadas = [], []
    for q, muestra in por_pregunta.items():
        if q in EXCLUIDAS_MUESTREO:
            bloqueadas.append(f'{q}: excluida del muestreo hasta terminar su corrección.')
            continue
        ok = all(m['aprobado'].lower() in SI and not m['comentario'] and m['reviso'] in REVISORES for m in muestreo_validas(muestra, datos))
        if not ok or len(muestreo_validas(muestra, datos)) != len(muestra):
            bloqueadas.append(f'{q}: el muestreo tiene filas sin aprobar, con "no" o con comentario.')
            continue
        quien = '/'.join(sorted({m['reviso'] for m in muestra}))
        for f in filas:
            if f['hoja'] != 'Resto' or f['pregunta'] != q or f['aprobado'] or f['comentario']:
                continue
            d = datos.get(f['iso'])
            r = d['respuestas'].get(q) if d else None
            if r is None or r.get('estado') in APROBADOS:
                continue
            if f['huella'] and f['huella'] != huella(r):
                revisar.append((f, 'No se aprobó por muestreo: la respuesta cambió desde que se generó la planilla.'))
                continue
            r['estado'] = 'aprobada_por_muestreo'
            r['revisado_por'] = quien
            r['revisado_fecha'] = hoy
            r['revision_origen'] = f'{base}: muestreo de {q} aprobado ({len(muestra)} filas)'
            d['_cambio'] = True
            aprobadas.append(f"{f['iso']} {q} ({quien})")
    return aprobadas, bloqueadas


def muestreo_validas(muestra, datos):
    """Filas del muestreo cuya respuesta existe y no cambió desde la planilla."""
    out = []
    for m in muestra:
        d = datos.get(m['iso'])
        r = d['respuestas'].get(m['pregunta']) if d else None
        if r is not None and (not m['huella'] or m['huella'] == huella(r)):
            out.append(m)
    return out


def informe(base, hoy, aprobadas, ya, corregir, revisar, muestreo=(), bloqueadas=()):
    def c(v):
        return (v or '—').replace('|', '\\|').replace('\n', ' ')
    partes = [f'# Informe de la planilla completada {base}', '', '**Confidencial: no se publica.**', '',
              f'Procesada el {hoy} con `python3 scripts/leer_completadas.py`.', '',
              f'- Aprobadas ahora: {len(aprobadas)}', f'- Ya estaban aprobadas: {len(ya)}',
              f'- Aprobadas por muestreo: {len(muestreo)}',
              f'- Para corregir (con "no" o comentario): {len(corregir)}', f'- Para revisar (datos incompletos o respuesta cambiada): {len(revisar)}', '']
    if bloqueadas:
        partes += ['## Preguntas sin aprobación por muestreo', ''] + [f'- {b}' for b in bloqueadas] + ['']
    if corregir:
        partes += ['## Para corregir', '', 'Corregir según el comentario y confirmar con quien revisó.', '',
                   '| País | # | Respuesta actual | Aprobado | Revisó | Comentario |', '|---|---|---|---|---|---|']
        partes += [f"| {f['iso']} | {f['pregunta']} | {c(str(r.get('respuesta')))} | {c(f['aprobado'])} | {c(f['reviso'])} | {c(f['comentario'])} |" for f, r in corregir]
        partes.append('')
    if revisar:
        partes += ['## Para revisar', '', '| País | # | Hoja y fila | Motivo |', '|---|---|---|---|']
        partes += [f"| {f['iso']} | {f['pregunta']} | {c(f['hoja'])}, fila {f['fila']} | {c(m)} |" for f, m in revisar]
        partes.append('')
    if aprobadas:
        partes += ['## Aprobadas ahora', '', ', '.join(aprobadas), '']
    if muestreo:
        partes += ['## Aprobadas por muestreo', '', ', '.join(muestreo), '']
    return '\n'.join(partes)


def main():
    hoy = datetime.date.today().isoformat()
    datos = {}
    for p in glob.glob(os.path.join(RESPUESTAS, '*.json')):
        d = json.load(open(p, encoding='utf-8'))
        d['_ruta'] = p
        datos[d['pais']] = d
    nombres = {d['nombre']: iso for iso, d in datos.items()}
    archivos = sorted(glob.glob(os.path.join(COMPLETADAS, '*.xlsx')) + glob.glob(os.path.join(COMPLETADAS, '*.md')))
    archivos = [a for a in archivos if not os.path.basename(a).startswith(('informe_', 'LEEME', '~$'))]
    reaplicar = '--reaplicar' in sys.argv
    if reaplicar:
        archivos = sorted(glob.glob(os.path.join(PROCESADAS, '*.xlsx'))) + archivos
    if not archivos:
        print('No hay planillas nuevas en data/revision/completadas/.')
        return
    os.makedirs(PROCESADAS, exist_ok=True)
    for ruta in archivos:
        base = os.path.basename(ruta)
        aprobadas, ya, corregir, revisar, muestreo, bloqueadas = procesar(ruta, datos, nombres, hoy)
        sufijo = '_reaplicada' if os.path.dirname(ruta) == PROCESADAS else ''
        nombre_inf = f"informe_{os.path.splitext(base)[0]}{sufijo}_{hoy}.md"
        with open(os.path.join(COMPLETADAS, nombre_inf), 'w', encoding='utf-8') as fh:
            fh.write(informe(base, hoy, aprobadas, ya, corregir, revisar, muestreo, bloqueadas) + '\n')
        if os.path.dirname(ruta) != PROCESADAS:
            shutil.move(ruta, os.path.join(PROCESADAS, base))
        print(f'{base}: {len(aprobadas)} aprobadas, {len(ya)} ya aprobadas, {len(muestreo)} por muestreo, {len(corregir)} para corregir, {len(revisar)} para revisar. Informe: {nombre_inf}')
    for d in datos.values():
        if d.pop('_cambio', False):
            ruta = d.pop('_ruta')
            with open(ruta, 'w', encoding='utf-8') as fh:
                json.dump(d, fh, ensure_ascii=False, indent=2)
                fh.write('\n')


if __name__ == '__main__':
    main()
